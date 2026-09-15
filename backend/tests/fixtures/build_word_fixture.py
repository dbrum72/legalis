"""Rebuild the neutral Word fixture used to check template formatting."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor

target = Path(__file__).parent / 'word-template.docx'
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(1)
section.left_margin = section.right_margin = Inches(1)
normal = doc.styles['Normal']
normal.font.name, normal.font.size = 'Calibri', Pt(12)
normal.font.color.rgb = RGBColor(0, 0, 0)
for name in ['Title', 'Heading 1']:
    doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
section.header.paragraphs[0].text = '{{ escritorio.nome }}'
doc.add_heading('Documento de demonstração', 0)
doc.add_paragraph('Modelo técnico para conferir a substituição de campos no Legalis.')
p = doc.add_paragraph('Pasta: ')
p.add_run('{{ pas').bold = True
p.add_run('ta.nome }}').bold = True
doc.add_paragraph('Data: {{ data.hoje }}')
doc.add_heading('Texto informado', 1)
doc.add_paragraph('{{ campo.texto }}')
table = doc.add_table(rows=1, cols=2)
table.style = 'Table Grid'
table.rows[0].cells[0].text = 'Campo'
table.rows[0].cells[1].text = 'Valor preenchido'
row = table.add_row()
row.cells[0].text = 'Referência'
row.cells[1].text = '{{ campo.referencia }}'
section.footer.paragraphs[0].text = 'Demonstração técnica — {{ data.hoje }}'
doc.save(target)
print(target)
